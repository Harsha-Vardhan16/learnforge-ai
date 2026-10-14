import re
from pathlib import Path
from typing import List, Dict, Any
from pptx import Presentation

def extract_text_from_pptx(file_path: str, source_name: str = None) -> List[Dict[str, Any]]:
    """
    Extracts text slide-by-slide from a PPTX file using python-pptx.
    Captures slide titles, body shapes, tables, and speaker notes.
    """
    path = Path(file_path)
    if not path.exists():
        raise FileNotFoundError(f"Presentation file not found at: {file_path}")

    if not source_name:
        source_name = path.stem.replace("_", " ").title()

    prs = Presentation(str(path))
    slides_data: List[Dict[str, Any]] = []

    for idx, slide in enumerate(prs.slides):
        slide_texts: List[str] = []

        # Extract text from all shapes
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    line = paragraph.text.strip()
                    if line:
                        slide_texts.append(line)
            elif shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        text = cell.text.strip()
                        if text:
                            slide_texts.append(text)

        # Extract speaker notes if present
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                slide_texts.append(f"[Notes: {notes}]")

        combined = " ".join(slide_texts)
        cleaned = re.sub(r'\s+', ' ', combined).strip()

        if cleaned:
            slides_data.append({
                "source_type": "slides",
                "source_name": source_name,
                "slide": idx + 1, # 1-indexed for slide citations
                "text": cleaned,
                "file_path": str(path)
            })

    return slides_data
